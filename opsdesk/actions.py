import fcntl
import json
import os
import smtplib
import tempfile
from datetime import timedelta
from email.message import EmailMessage
from pathlib import Path

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from opsdesk.storage import DeliveryReceipt, SupportRequest, WorkOrder, new_id, timestamp


class ActionHandlers:
    """Tool implementations accept persisted operation IDs, never arbitrary commands."""

    def __init__(self, database, config):
        self.database, self.config = database, config

    def create_work_order(self, request_id):
        with self.database.session.begin() as db:
            request = db.scalar(select(SupportRequest).where(SupportRequest.id == request_id).with_for_update())
            if request is None or request.phase != "COMPLETE" or request.outcome["intake"]["intent"] == "FAQ":
                raise ValueError("Operation is not eligible for a work order")
            existing = db.scalar(select(WorkOrder).where(WorkOrder.request_id == request_id))
            if existing:
                return {"reference": existing.reference, "reused": True}
            order = WorkOrder(id=new_id(), reference="IT-" + request_id[:12].upper(), request_id=request_id)
            db.add(order)
            return {"reference": order.reference, "reused": False}

    def export_ledger(self, request_id):
        with self.database.session() as db:
            request = db.get(SupportRequest, request_id)
            order = db.scalar(select(WorkOrder).where(WorkOrder.request_id == request_id))
            if not order:
                raise ValueError("Work order must be created first")
            row = [request_id, order.reference, request.created.isoformat(), request.outcome["intake"]["summary"], request.outcome["priority"]["level"], order.phase]
        path = Path(self.config.ledger_file)
        path.parent.mkdir(parents=True, exist_ok=True)
        # The sidecar lock survives atomic replacement of the workbook.
        with path.with_suffix(".lock").open("a") as lock:
            fcntl.flock(lock.fileno(), fcntl.LOCK_EX)
            workbook = load_workbook(path) if path.exists() else Workbook()
            sheet = workbook.active
            if not path.exists():
                sheet.title = "Requests"
                sheet.append(["operation_id", "work_order", "created_utc", "summary", "priority", "status_at_export"])
            existing = any(values[0] == request_id for values in sheet.iter_rows(min_row=2, values_only=True))
            if not existing:
                sheet.append(row)
                fd, temporary = tempfile.mkstemp(suffix=".xlsx", dir=path.parent)
                os.close(fd)
                try:
                    workbook.save(temporary)
                    os.replace(temporary, path)
                finally:
                    Path(temporary).unlink(missing_ok=True)
            workbook.close()
        self._receipt(request_id, "ledger", {"file": path.name})
        return {"exported": True, "reused": existing, "file": path.name}

    def notify_escalation(self, request_id):
        with self.database.session.begin() as db:
            request = db.scalar(select(SupportRequest).where(SupportRequest.id == request_id).with_for_update())
            if request is None or request.outcome["priority"]["level"] != "P1":
                raise ValueError("Only P1 operations can be escalated")
            existing = db.scalar(select(DeliveryReceipt).where(DeliveryReceipt.request_id == request_id, DeliveryReceipt.kind == "escalation"))
            if existing:
                return {**existing.detail, "reused": True}
            count = db.scalar(select(func.count()).select_from(DeliveryReceipt).where(DeliveryReceipt.kind == "escalation", DeliveryReceipt.created > timestamp() - timedelta(hours=1)))
            if count >= self.config.notification_limit:
                raise ValueError("Hourly escalation limit reached")
            mode = self.config.notification_mode
            if mode == "smtp":
                if not all([self.config.smtp_host, self.config.smtp_sender, self.config.smtp_recipient]):
                    raise ValueError("SMTP configuration is incomplete")
                email = EmailMessage()
                email["From"], email["To"] = self.config.smtp_sender, self.config.smtp_recipient
                email["Subject"] = "OpsDesk P1 incident " + request_id[:12]
                email.set_content(json.dumps({"operation": request_id, "summary": request.outcome["intake"]["summary"]}, ensure_ascii=False))
                with smtplib.SMTP(self.config.smtp_host, self.config.smtp_port, timeout=15) as sender:
                    sender.starttls()
                    if self.config.smtp_username:
                        sender.login(self.config.smtp_username, self.config.smtp_password)
                    sender.send_message(email)
            elif mode != "record":
                raise ValueError("Unsupported notification mode")
            detail = {"mode": mode, "status": "sent" if mode == "smtp" else "recorded"}
            db.add(DeliveryReceipt(request_id=request_id, kind="escalation", detail=detail))
            return {**detail, "reused": False}

    def _receipt(self, request_id, kind, detail):
        with self.database.session.begin() as db:
            db.scalar(select(SupportRequest).where(SupportRequest.id == request_id).with_for_update())
            if not db.scalar(select(DeliveryReceipt).where(DeliveryReceipt.request_id == request_id, DeliveryReceipt.kind == kind)):
                db.add(DeliveryReceipt(request_id=request_id, kind=kind, detail=detail))

    def dispatch(self, kind, request_id):
        handlers = {"create_work_order": self.create_work_order, "export_ledger": self.export_ledger, "notify_escalation": self.notify_escalation}
        if kind not in handlers:
            raise ValueError("Tool is not allowlisted")
        return handlers[kind](request_id)
