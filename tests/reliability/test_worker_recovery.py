import asyncio
import sqlite3
import threading

import httpx
import pytest
from openpyxl import load_workbook
from pymysql.err import OperationalError as MySQLError
from sqlalchemy.exc import IntegrityError, OperationalError, ProgrammingError, TimeoutError as PoolTimeout

from opsdesk import queueing
from opsdesk.main import create_app
from opsdesk.queueing import ActionWorker
from opsdesk.storage import QueuedAction
from tests.reliability.helpers import Clock, operation, snapshot, stop_after_terminal


def db_error(kind):
    if kind in {"busy", "schema"}:
        original = sqlite3.OperationalError("synthetic-private-error-text")
        original.sqlite_errorcode = sqlite3.SQLITE_BUSY if kind == "busy" else sqlite3.SQLITE_ERROR
        return OperationalError("synthetic", None, original)
    if kind in {"deadlock", "disconnect", "auth"}:
        code = {"deadlock": 1213, "disconnect": 2003, "auth": 1045}[kind]
        return OperationalError("synthetic", None, MySQLError(code, "synthetic-private-error-text"))
    if kind == "pool":
        return PoolTimeout("synthetic-private-error-text")
    if kind == "constraint":
        return IntegrityError("synthetic", None, ValueError("synthetic-private-error-text"))
    if kind == "programming":
        return ProgrammingError("synthetic", None, ValueError("synthetic-private-error-text"))
    return RuntimeError("synthetic-private-error-text")


@pytest.mark.parametrize("kind", ["busy", "deadlock", "disconnect", "pool"])
def test_claim_database_recovers_and_completes_real_actions(runtime, monkeypatch, record, kind):
    config, db, _, _, owner, _ = runtime
    operation(db, owner)
    worker = ActionWorker(db, config)
    stop_after_terminal(worker, db, monkeypatch)
    claim, delays, attempts = worker.claim, [], []

    def flaky():
        attempts.append(True)
        if len(attempts) <= 2:
            raise db_error(kind)
        return claim()

    async def pause(delay):
        delays.append(delay)
        await asyncio.sleep(0)

    monkeypatch.setattr(worker, "claim", flaky)
    monkeypatch.setattr(worker, "_pause", pause, raising=False)

    async def scenario():
        try:
            await asyncio.wait_for(worker.run(), timeout=2)
        finally:
            record({"fault": kind, "delays": delays, "after": snapshot(db), "worker": getattr(worker, "observations", None)})
    asyncio.run(scenario())
    assert delays == [.4, .8]
    state = snapshot(db)
    assert all(a["phase"] == "DONE" and a["tries"] == 1 and not a["leased"] for a in state["actions"])
    assert state["work_orders"] == state["receipts"] == 1
    assert [a["result"] for a in state["audits"]] == ["DONE", "DONE"]
    assert worker.observations["recoveries"] == 1 and worker.observations["consecutive_errors"] == 0
    workbook = load_workbook(config.ledger_file)
    assert workbook.active.max_row == 2
    workbook.close()


def test_repeated_transient_errors_use_capped_backoff(runtime, monkeypatch, record):
    config, db, _, _, owner, _ = runtime
    operation(db, owner)
    worker = ActionWorker(db, config)
    stop_after_terminal(worker, db, monkeypatch)
    real_claim, delays, calls = worker.claim, [], 0

    def flaky():
        nonlocal calls
        calls += 1
        if calls <= 9:
            raise db_error("busy")
        return real_claim()

    async def pause(delay):
        delays.append(delay)
        await asyncio.sleep(0)

    monkeypatch.setattr(worker, "claim", flaky)
    monkeypatch.setattr(worker, "_pause", pause, raising=False)
    try:
        asyncio.run(asyncio.wait_for(worker.run(), 2))
    finally:
        record({"fault_count": 9, "scheduled_delays_not_wall_time": delays, "after": snapshot(db)})
    assert delays == [.4, .8, 1.6, 3.2, 6.4, 12.8, 25.6, 30, 30]
    assert all(a["phase"] == "DONE" for a in snapshot(db)["actions"])


@pytest.mark.parametrize("kind", ["schema", "auth", "constraint", "programming", "unknown"])
def test_permanent_database_or_programming_error_stops_without_retry(runtime, monkeypatch, record, caplog, kind):
    config, db, _, _, owner, _ = runtime
    operation(db, owner)
    worker = ActionWorker(db, config)
    error, calls, delays = db_error(kind), [], []

    def broken():
        calls.append(True)
        raise error

    async def pause(delay):
        delays.append(delay)

    monkeypatch.setattr(worker, "claim", broken)
    monkeypatch.setattr(worker, "_pause", pause, raising=False)
    with pytest.raises(type(error)):
        asyncio.run(worker.run())
    record({"fault": kind, "claim_calls": len(calls), "delays": delays, "after": snapshot(db), "worker": getattr(worker, "observations", None)})
    assert len(calls) == 1 and not delays
    assert worker.observations["state"] == "failed"
    assert worker.observations["last_error"] == type(error).__name__
    assert "worker" in caplog.text.lower() and "synthetic-private-error-text" not in caplog.text
    assert all(a["phase"] == "READY" and a["tries"] == 0 for a in snapshot(db)["actions"])


@pytest.mark.parametrize("when", ["before-commit", "after-commit"])
def test_finish_database_fault_recovers_without_duplicate_side_effects(runtime, monkeypatch, record, when):
    config, db, _, _, owner, _ = runtime
    operation(db, owner)
    clock = Clock()
    monkeypatch.setattr(queueing, "timestamp", clock)
    worker = ActionWorker(db, config)
    finish, invoke, injected, delays, dispatched = worker.finish, worker.invoke, False, [], []

    async def observed_invoke(action):
        result = await invoke(action)
        dispatched.append({"kind": action["kind"], "reused": result.get("reused")})
        return result

    def flaky_finish(action, **kwargs):
        nonlocal injected
        if action["kind"] == "create_work_order" and not injected:
            injected = True
            if when == "after-commit":
                finish(action, **kwargs)
            raise db_error("disconnect")
        result = finish(action, **kwargs)
        if all(a["phase"] == "DONE" for a in snapshot(db)["actions"]):
            worker.stopping = True
        return result

    async def pause(delay):
        delays.append(delay)
        clock.advance(91)
        await asyncio.sleep(0)

    monkeypatch.setattr(worker, "invoke", observed_invoke)
    monkeypatch.setattr(worker, "finish", flaky_finish)
    monkeypatch.setattr(worker, "_pause", pause, raising=False)
    try:
        asyncio.run(asyncio.wait_for(worker.run(), 2))
    finally:
        record({"fault": when, "clock": "simulated lease expiry", "delays": delays, "dispatches": dispatched, "after": snapshot(db)})
    assert delays == [.4]
    state = snapshot(db)
    assert all(a["phase"] == "DONE" for a in state["actions"])
    assert state["work_orders"] == state["receipts"] == 1
    assert "FAILED" not in [a["result"] for a in state["audits"]]
    creations = [a for a in dispatched if a["kind"] == "create_work_order"]
    assert [a["reused"] for a in creations] == ([False, True] if when == "before-commit" else [False])
    assert sum(a["result"] == "LEASE_EXPIRED" for a in state["audits"]) == int(when == "before-commit")


def test_cancel_during_backoff_then_restart_completes_pending_queue(runtime, monkeypatch, record):
    config, db, _, _, owner, _ = runtime
    operation(db, owner)
    worker = ActionWorker(db, config)

    async def scenario():
        entered = asyncio.Event()

        def broken():
            raise db_error("busy")

        async def pause(delay):
            entered.set()
            await asyncio.Event().wait()

        monkeypatch.setattr(worker, "claim", broken)
        monkeypatch.setattr(worker, "_pause", pause, raising=False)
        job = asyncio.create_task(worker.run())
        try:
            await asyncio.wait_for(entered.wait(), .5)
            job.cancel()
            with pytest.raises(asyncio.CancelledError):
                await job
            assert worker.observations["state"] == "stopped"
            before = snapshot(db)
            assert all(a["phase"] == "READY" and a["tries"] == 0 for a in before["actions"])
            restarted = ActionWorker(db, config)
            stop_after_terminal(restarted, db, monkeypatch)
            await asyncio.wait_for(restarted.run(), 2)
            assert all(a["phase"] == "DONE" for a in snapshot(db)["actions"])
            record({"cancelled_during": "backoff", "before_restart": before, "after_restart": snapshot(db)})
        finally:
            job.cancel()
            await asyncio.gather(job, return_exceptions=True)
    asyncio.run(scenario())


def test_cancel_inflight_preserves_lease_and_restarts_after_expiry(runtime, monkeypatch, record):
    config, db, _, _, owner, _ = runtime
    operation(db, owner)
    clock = Clock()
    monkeypatch.setattr(queueing, "timestamp", clock)
    worker = ActionWorker(db, config)

    async def scenario():
        entered, exited = asyncio.Event(), asyncio.Event()

        async def blocked(action):
            entered.set()
            try:
                await asyncio.Event().wait()
            finally:
                exited.set()

        monkeypatch.setattr(worker, "invoke", blocked)
        job = asyncio.create_task(worker.run())
        try:
            await asyncio.wait_for(entered.wait(), .5)
            job.cancel()
            with pytest.raises(asyncio.CancelledError):
                await job
            assert exited.is_set()
            before = snapshot(db)
            with db.session() as session:
                row = next(a for a in session.query(QueuedAction) if a.phase == "RUNNING")
                stale = {"id": row.id, "lease": row.lease}
            restarted = ActionWorker(db, config)
            assert restarted.claim() is None
            clock.advance(91)
            assert await restarted.step() and await restarted.step()
            assert not worker.finish(stale, result={"stale": True})
            state = snapshot(db)
            assert state["work_orders"] == state["receipts"] == 1
            assert all(a["phase"] == "DONE" for a in state["actions"])
            assert sum(a["result"] == "LEASE_EXPIRED" for a in state["audits"]) == 1
            record({"cancelled_during": "async invocation", "clock": "simulated", "before_restart": before, "after_restart": state})
        finally:
            job.cancel()
            await asyncio.gather(job, return_exceptions=True)
    asyncio.run(scenario())


def test_cancel_waits_for_owned_direct_thread_before_restarting(runtime, monkeypatch, record):
    config, db, _, _, owner, _ = runtime
    operation(db, owner)
    worker = ActionWorker(db, config)
    dispatch = worker.handlers.dispatch
    started, release, finished = threading.Event(), threading.Event(), threading.Event()

    def blocked(kind, request_id):
        result = dispatch(kind, request_id)
        if kind == "create_work_order":
            started.set()
            try:
                assert release.wait(2), "Test must release its own thread"
            finally:
                finished.set()
        return result

    monkeypatch.setattr(worker.handlers, "dispatch", blocked)

    async def scenario():
        job = asyncio.create_task(worker.run())
        async def unblock():
            await asyncio.sleep(.05)
            release.set()
        unblock_job = None
        try:
            assert await asyncio.to_thread(started.wait, 1)
            unblock_job = asyncio.create_task(unblock())
            job.cancel()
            with pytest.raises(asyncio.CancelledError):
                await job
            assert finished.is_set(), "Cancelled worker must not leave its direct thread unowned"
            before = snapshot(db)
            assert before["work_orders"] == 1 and any(a["phase"] == "RUNNING" for a in before["actions"])
            clock = Clock()
            clock.advance(91)
            monkeypatch.setattr(queueing, "timestamp", clock)
            restarted = ActionWorker(db, config)
            assert await restarted.step() and await restarted.step()
            state = snapshot(db)
            assert state["work_orders"] == state["receipts"] == 1
            assert all(a["phase"] == "DONE" for a in state["actions"])
            record({"cancelled_during": "real local direct thread", "thread_joined_before_return": True, "before_restart": before, "after_restart": state})
        finally:
            release.set()
            await asyncio.to_thread(finished.wait, 2)
            job.cancel()
            await asyncio.gather(job, return_exceptions=True)
            if unblock_job:
                unblock_job.cancel()
                await asyncio.gather(unblock_job, return_exceptions=True)
    asyncio.run(scenario())


@pytest.mark.parametrize("body_error", [False, True], ids=["normal-close", "exceptional-close"])
def test_app_lifespan_cleans_worker_and_database(runtime, monkeypatch, record, body_error):
    config = runtime[0]
    config.worker_enabled = True
    app = create_app(config)
    worker = app.state.worker
    dispose, disposed = app.state.database.engine.dispose, []
    monkeypatch.setattr(app.state.database.engine, "dispose", lambda: (disposed.append(True), dispose())[-1])

    async def scenario():
        entered, exited = asyncio.Event(), asyncio.Event()
        captured = []

        async def held():
            captured.append(asyncio.current_task())
            entered.set()
            try:
                await asyncio.Event().wait()
            finally:
                exited.set()

        monkeypatch.setattr(worker, "run", held)
        try:
            try:
                async with app.router.lifespan_context(app):
                    await asyncio.wait_for(entered.wait(), .5)
                    if body_error:
                        raise ValueError("Synthetic lifespan body failure")
            except* ValueError as errors:
                # MCP's AnyIO task group wraps the original body error. Accept
                # that documented shape without masking unrelated exceptions.
                assert body_error
                pending = list(errors.exceptions)
                while pending:
                    error = pending.pop()
                    if isinstance(error, BaseExceptionGroup):
                        pending.extend(error.exceptions)
                    else:
                        assert str(error) == "Synthetic lifespan body failure"
            assert exited.is_set() and captured[0].done()
            assert disposed
            record({"body_error": body_error, "worker_done": captured[0].done(), "database_disposed": bool(disposed)})
        finally:
            for task in captured:
                task.cancel()
            await asyncio.gather(*captured, return_exceptions=True)
            dispose()
    asyncio.run(scenario())


def test_fatal_background_worker_is_visible_then_fresh_app_recovers(runtime, monkeypatch, record):
    config, db, _, _, owner, _ = runtime
    operation(db, owner)
    config.worker_enabled = True
    app = create_app(config)
    worker = app.state.worker
    monkeypatch.setattr(worker, "claim", lambda: (_ for _ in ()).throw(db_error("schema")))

    async def scenario():
        async with app.router.lifespan_context(app):
            # Await the owned task if available, without suppressing the final assertions.
            job = getattr(app.state, "worker_task", None)
            if job:
                await asyncio.gather(job, return_exceptions=True)
            else:
                await asyncio.sleep(0)
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://testserver") as client:
                health = await client.get("/api/health")
                assert health.status_code == 503
                assert health.json()["worker"]["state"] == "failed"
                await client.post("/api/login", json={"username": "engineer", "password": config.engineer_password})
                diagnostics = (await client.get("/api/diagnostics")).json()
                assert diagnostics["worker"]["alive"] is False
                assert diagnostics["worker"]["last_error"] == "OperationalError"
                record({"fatal_health": health.json(), "before_restart": snapshot(db)})
        fresh = create_app(config)
        stop_after_terminal(fresh.state.worker, fresh.state.database, monkeypatch)
        async with fresh.router.lifespan_context(fresh):
            await asyncio.wait_for(fresh.state.worker_task, 2)
            state = snapshot(fresh.state.database)
            assert all(a["phase"] == "DONE" for a in state["actions"])
            assert state["work_orders"] == state["receipts"] == 1
            record({"fresh_app_completed_pending_actions": state})
    asyncio.run(scenario())
