import asyncio
from datetime import timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import event, select, update

from opsdesk import queueing
from opsdesk.actions import ActionHandlers
from opsdesk.queueing import ActionWorker
from opsdesk.storage import QueuedAction, timestamp
from tests.reliability.helpers import Clock, operation, snapshot


def test_sixty_task_blocked_window_drains_in_dependency_order_without_duplicates(runtime, record):
    config, db, _, _, owner, _ = runtime
    requests = [operation(db, owner) for _ in range(30)]
    now = timestamp()
    with db.session.begin() as session:
        session.execute(update(QueuedAction).where(QueuedAction.kind == "create_work_order").values(due=now - timedelta(seconds=1)))
        session.execute(update(QueuedAction).where(QueuedAction.kind == "export_ledger").values(due=now - timedelta(seconds=10)))
    worker = ActionWorker(db, config)
    before = snapshot(db)
    statements = []
    def observed(*args):
        statements.append(args[2])
    event.listen(db.engine, "before_cursor_execute", observed)
    try:
        first = worker.claim()
    finally:
        event.remove(db.engine, "before_cursor_execute", observed)
        record({"original_audit_fixture": {"ready": 60, "parents": 30, "older_blocked_children": 30}, "first_claim": first, "claim_sql_count": len(statements), "before": before})
    assert first is not None and first["kind"] == "create_work_order"
    assert len(statements) <= 8, "Claim should not query each blocked child's parent separately"

    async def drain():
        worker.finish(first, result=await worker.invoke(first))
        for _ in range(59):
            assert await worker.step()
        assert not await worker.step()
    asyncio.run(drain())
    state = snapshot(db)
    assert len(state["actions"]) == 60 and all(a["phase"] == "DONE" and a["tries"] == 1 for a in state["actions"])
    assert state["work_orders"] == state["receipts"] == 30
    assert len(state["audits"]) == 60 and all(a["result"] == "DONE" for a in state["audits"])
    tools = ActionHandlers(db, config)
    assert tools.create_work_order(requests[0])["reused"] and tools.export_ledger(requests[0])["reused"]
    workbook = load_workbook(config.ledger_file)
    assert workbook.active.max_row == 31
    workbook.close()
    assert snapshot(db)["work_orders"] == 30 and snapshot(db)["receipts"] == 30
    record({"after_draining": state, "ledger_data_rows": 30, "replay_created_no_extra_rows": True})


def test_dead_dependency_cleanup_does_not_block_independent_runnable_action(runtime, record):
    config, db, _, _, owner, _ = runtime
    for _ in range(30):
        operation(db, owner, "P1")
    now = timestamp()
    with db.session.begin() as session:
        session.execute(update(QueuedAction).where(QueuedAction.kind == "create_work_order").values(phase="DEAD", tries=config.max_action_attempts))
        session.execute(update(QueuedAction).where(QueuedAction.depends_on.is_not(None)).values(due=now - timedelta(seconds=10)))
    independent = operation(db, owner)
    worker = ActionWorker(db, config)
    first = worker.claim()
    record({"dead_parents": 30, "blocked_children": 60, "first_claim": first})
    assert first and first["request_id"] == independent and first["kind"] == "create_work_order"
    async def scenario():
        worker.finish(first, result=await worker.invoke(first))
        assert await worker.step()
        for _ in range(3):
            assert not await worker.step()
    asyncio.run(scenario())
    state = snapshot(db)
    assert sum(a["phase"] == "DEAD" for a in state["actions"]) == 90
    assert sum(a["phase"] == "DONE" for a in state["actions"]) == 2
    assert state["work_orders"] == state["receipts"] == 1
    record({"after": state})


@pytest.mark.parametrize("parent_phase", ["READY", "RUNNING", "DONE", "DEAD"])
def test_child_requires_done_parent_and_dead_parent_blocks_child(runtime, record, parent_phase):
    config, db, _, _, owner, _ = runtime
    request_id = operation(db, owner)
    now = timestamp()
    with db.session.begin() as session:
        parent = session.scalar(select(QueuedAction).where(QueuedAction.kind == "create_work_order"))
        parent.phase, parent.due = parent_phase, now + timedelta(hours=1)
        if parent_phase == "RUNNING":
            parent.lease, parent.lease_until, parent.tries = "synthetic", now + timedelta(hours=1), 1
    worker = ActionWorker(db, config)
    if parent_phase == "DONE":
        worker.handlers.create_work_order(request_id)
    claimed = worker.claim()
    if parent_phase == "DONE":
        assert claimed and claimed["kind"] == "export_ledger"
        worker.finish(claimed, result=asyncio.run(worker.invoke(claimed)))
        assert snapshot(db)["receipts"] == 1
    else:
        assert claimed is None
        child = next(a for a in snapshot(db)["actions"] if a["kind"] == "export_ledger")
        assert child["phase"] == ("DEAD" if parent_phase == "DEAD" else "READY")
        assert child["tries"] == 0 and not child["leased"]
    record({"parent_phase": parent_phase, "claimed": claimed, "after": snapshot(db)})


def test_empty_or_future_due_queue_is_not_claimed(runtime, record):
    config, db, _, _, owner, _ = runtime
    worker = ActionWorker(db, config)
    assert worker.claim() is None
    operation(db, owner)
    with db.session.begin() as session:
        session.execute(update(QueuedAction).values(due=timestamp() + timedelta(hours=1)))
    assert worker.claim() is None
    state = snapshot(db)
    assert all(a["phase"] == "READY" and a["tries"] == 0 and not a["leased"] for a in state["actions"])
    record({"future_due_after": state})


def test_exhausted_ready_action_never_gets_an_extra_attempt(runtime, record):
    config, db, _, _, owner, _ = runtime
    operation(db, owner)
    with db.session.begin() as session:
        session.scalar(select(QueuedAction).where(QueuedAction.kind == "create_work_order")).tries = config.max_action_attempts
    worker = ActionWorker(db, config)
    claimed = worker.claim()
    state = snapshot(db)
    record({"claimed": claimed, "after": state})
    assert claimed is None
    assert all(a["phase"] == "DEAD" for a in state["actions"])
    assert next(a for a in state["actions"] if a["kind"] == "create_work_order")["tries"] == config.max_action_attempts


@pytest.mark.parametrize("exhausted", [False, True], ids=["last-allowed-attempt", "budget-exhausted"])
def test_expired_lease_respects_budget_and_dependency(runtime, monkeypatch, record, exhausted):
    config, db, _, _, owner, _ = runtime
    operation(db, owner)
    clock = Clock()
    monkeypatch.setattr(queueing, "timestamp", clock)
    with db.session.begin() as session:
        parent = session.scalar(select(QueuedAction).where(QueuedAction.kind == "create_work_order"))
        parent.phase, parent.lease, parent.lease_until = "RUNNING", "expired", clock() - timedelta(seconds=1)
        parent.tries = config.max_action_attempts if exhausted else config.max_action_attempts - 1
        old = {"id": parent.id, "lease": parent.lease}
    worker = ActionWorker(db, config)
    claim = worker.claim()
    if exhausted:
        assert claim is None and all(a["phase"] == "DEAD" for a in snapshot(db)["actions"])
    else:
        assert claim and claim["kind"] == "create_work_order" and claim["lease"] != "expired"
        worker.finish(claim, result=asyncio.run(worker.invoke(claim)))
        assert asyncio.run(worker.step())
        assert all(a["phase"] == "DONE" for a in snapshot(db)["actions"])
    assert not worker.finish(old, result={"stale": True})
    state = snapshot(db)
    assert sum(a["result"] == "LEASE_EXPIRED" for a in state["audits"]) == 1
    assert next(a for a in state["actions"] if a["kind"] == "create_work_order")["tries"] == config.max_action_attempts
    record({"clock": "simulated", "exhausted": exhausted, "after": state})


def test_two_serial_workers_claim_distinct_actions_and_stale_lease_is_fenced(runtime, record):
    config, db, _, _, owner, _ = runtime
    operation(db, owner)
    operation(db, owner)
    first_worker, second_worker = ActionWorker(db, config), ActionWorker(db, config)
    first, second = first_worker.claim(), second_worker.claim()
    assert first and second and first["id"] != second["id"]
    assert not first_worker.finish({"id": first["id"], "lease": "wrong"}, result={})
    async def scenario():
        first_worker.finish(first, result=await first_worker.invoke(first))
        second_worker.finish(second, result=await second_worker.invoke(second))
        assert await first_worker.step() and await second_worker.step()
    asyncio.run(scenario())
    state = snapshot(db)
    assert all(a["phase"] == "DONE" and a["tries"] == 1 for a in state["actions"])
    assert state["work_orders"] == state["receipts"] == 2
    record({"execution": "serial SQLite claims, not MySQL concurrency", "claims_distinct": True, "after": state})


def test_action_failures_dead_letter_then_manual_dependency_recovery(runtime, monkeypatch, record):
    config, db, _, _, owner, _ = runtime
    operation(db, owner)
    clock = Clock()
    monkeypatch.setattr(queueing, "timestamp", clock)
    worker = ActionWorker(db, config)
    actual_invoke, calls = worker.invoke, []
    async def broken(action):
        calls.append(action["kind"])
        raise ValueError("Synthetic tool rejection")
    monkeypatch.setattr(worker, "invoke", broken)
    for _ in range(config.max_action_attempts):
        assert asyncio.run(worker.step())
        clock.advance(61)
    assert worker.claim() is None
    failed = snapshot(db)
    assert all(a["phase"] == "DEAD" for a in failed["actions"])
    assert calls == ["create_work_order"] * config.max_action_attempts
    assert len(failed["audits"]) == config.max_action_attempts and all(a["result"] == "FAILED" for a in failed["audits"])
    parent = next(a for a in failed["actions"] if a["kind"] == "create_work_order")
    child = next(a for a in failed["actions"] if a["kind"] == "export_ledger")
    with pytest.raises(ValueError):
        worker.retry(child["id"])
    monkeypatch.setattr(worker, "invoke", actual_invoke)
    worker.retry(parent["id"])
    assert asyncio.run(worker.step())
    worker.retry(child["id"])
    assert asyncio.run(worker.step())
    state = snapshot(db)
    assert all(a["phase"] == "DONE" and a["tries"] == 1 for a in state["actions"])
    assert state["work_orders"] == state["receipts"] == 1
    record({"clock": "simulated", "dead_letter": failed, "after_manual_dependency_recovery": state, "manual_actor_audit": "not implemented"})
