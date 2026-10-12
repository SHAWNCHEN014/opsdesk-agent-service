"""QA-SEC-02: literal Excel cells, file integrity and ledger replay."""
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from opsdesk.actions import ActionHandlers
from opsdesk.storage import DeliveryReceipt, SupportRequest, Thread, WorkOrder, new_id


def persist_incident(db, owner, summary):
    operation, thread = new_id(), new_id()
    with db.session.begin() as session:
        session.add(Thread(id=thread, owner=owner))
        session.flush()
        session.add(SupportRequest(id=operation, owner=owner, thread_id=thread,
            client_key=operation, input_digest="synthetic", message=summary, phase="COMPLETE",
            outcome={"intake": {"intent": "INCIDENT", "summary": summary}, "priority": {"level": "P2"}}))
    return operation


@pytest.mark.parametrize("text", [
    '=HYPERLINK("https://example.invalid","电脑无法开机")',
    "+SUM(1,2)", "-1+2", "@SUM(1,2)",
    '\t=HYPERLINK("https://example.invalid","tab")', "\r\n=1+1", " =1+1",
    "#DIV/0!", "电脑无法开机 & <排查> \"说明\" 😀", "第一行\n第二行\t字段",
], ids=["equals", "plus", "minus", "at", "tab-prefix", "crlf-prefix", "space-prefix", "excel-error-token", "unicode-xml", "multiline"])
def test_exported_summary_is_literal_and_preserves_text(runtime, text):
    config, db, _, _, owner, _ = runtime
    operation = persist_incident(db, owner, text)
    tools = ActionHandlers(db, config)
    tools.create_work_order(operation)
    assert not tools.export_ledger(operation)["reused"]
    assert tools.export_ledger(operation)["reused"]
    workbook = load_workbook(config.ledger_file, data_only=False)
    try:
        sheet = workbook.active
        assert sheet.max_row == 2 and sheet.max_column == 6
        assert sheet.cell(2, 4).data_type == "s"
        assert sheet.cell(2, 4).value == text
        assert list(sheet.values)[0] == ("operation_id", "work_order", "created_utc", "summary", "priority", "status_at_export")
        assert sheet.cell(2, 1).value == operation
        assert sheet.cell(2, 5).value == "P2" and sheet.cell(2, 6).value == "OPEN"
    finally:
        workbook.close()
    with ZipFile(config.ledger_file) as package:
        assert package.testzip() is None
        assert "[Content_Types].xml" in package.namelist()
        xml = ET.fromstring(package.read("xl/worksheets/sheet1.xml"))
        assert not xml.findall(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}f")
        assert not any("externalLinks/" in name for name in package.namelist())
    with db.session() as session:
        assert session.scalar(select(func.count()).select_from(WorkOrder)) == 1
        assert session.scalar(select(func.count()).select_from(DeliveryReceipt)) == 1


def test_append_keeps_existing_row_and_work_order_status(runtime):
    config, db, _, _, owner, _ = runtime
    tools = ActionHandlers(db, config)
    operations = [persist_incident(db, owner, text) for text in ("原有普通工单", "=1+1")]
    for operation in operations:
        tools.create_work_order(operation)
    with db.session.begin() as session:
        session.scalar(select(WorkOrder).where(WorkOrder.request_id == operations[1])).phase = "INVESTIGATING"
    tools.export_ledger(operations[0])
    tools.export_ledger(operations[1])
    assert tools.export_ledger(operations[0])["reused"]
    workbook = load_workbook(config.ledger_file)
    try:
        assert workbook.active.max_row == 3
        assert workbook.active.cell(2, 4).value == "原有普通工单"
        assert workbook.active.cell(3, 4).value == "=1+1"
        assert workbook.active.cell(3, 4).data_type == "s"
        assert workbook.active.cell(3, 6).value == "INVESTIGATING"
    finally:
        workbook.close()


def test_carriage_returns_survive_a_later_append(runtime):
    config, db, _, _, owner, _ = runtime
    tools = ActionHandlers(db, config)
    for text in ("第一行\r\n第二行", "新工单"):
        operation = persist_incident(db, owner, text)
        tools.create_work_order(operation)
        tools.export_ledger(operation)
    workbook = load_workbook(config.ledger_file)
    try:
        assert workbook.active.cell(2, 4).value == "第一行\r\n第二行"
        assert workbook.active.cell(3, 4).value == "新工单"
    finally:
        workbook.close()
